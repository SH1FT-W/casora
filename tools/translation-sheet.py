#!/usr/bin/env python3
"""Translation review sheet for a Casora language: export to Excel, import the corrections back.

    uv run --python 3.14 --with openpyxl tools/translation-sheet.py export es [out.xlsx] [--since <commit>]
    uv run --python 3.14 --with openpyxl tools/translation-sheet.py import es reviewed.xlsx [--dry-run]

Export writes one row per text from the four translation files of that language (dashboard phrases, Studio
panel, dashboard keys, Home Assistant setup), with German and English beside the translation, a "Correct?"
choice and a "Correction" column. Important rows come first (short labels, patterns with placeholders).
Import takes every row marked "Wrong" (or with a correction), checks the placeholders against the source and
writes the correction into the JSON files; then run tools/build-i18n.py and tools/build-panel-i18n.py.
Works for any language that has the four files (copy them from en/de to start a new one).

--since <commit> exports only the rows that are new or changed since that commit (source or translation), as
an addendum to a sheet sent out earlier. Import is safe for old sheets: a correction is only written when the
row still has the source and translation the reviewer saw; otherwise it is reported as a conflict and the
current text stays. Rows not in the sheet and keys added since are left as they are.
"""
import json
import os
import re
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TR = os.path.join(ROOT, "custom_components/casora/translations")
SLOT = re.compile(r"\{\w+\}|%s|<[^>]+>")
COLS = ["Key", "Where", "Priority", "German", "English", "Translation", "Correct?", "Correction", "Comment"]
FONT = "Arial"


def load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def flat(obj, path=""):
    if isinstance(obj, dict):
        for k, v in obj.items():
            yield from flat(v, path + "/" + k)
    elif isinstance(obj, str):
        yield path, obj


def priority(text, pattern, dashboard):
    # 1 = what everyone sees daily: dashboard labels and patterns; 2 = Studio labels and short sentences;
    # 3 = long help texts.
    if dashboard and (pattern or len(text) <= 40):
        return 1
    return 2 if len(text) <= 90 else 3


def rows(lang, tr_dir=None):
    """[(key, where, prio, german, english, translation)] – key is what import matches on."""
    TR = tr_dir or globals()["TR"]
    out = []
    # Dashboard: German source, English table beside it.
    en = load(os.path.join(TR, "dashboard/phrases/en.json"))
    tr = load(os.path.join(TR, "dashboard/phrases/%s.json" % lang))
    for de_txt, en_txt in en["exact"].items():
        out.append(("dash:exact:" + de_txt, "Dashboard", priority(de_txt, False, True), de_txt, en_txt, tr["exact"].get(de_txt, "")))
    # English leftovers the dashboard translates as well ("Favorites", "Home"): German from phrases/de.json.
    de_path = os.path.join(TR, "dashboard/phrases/de.json")
    de_extra = load(de_path).get("exact", {}) if os.path.exists(de_path) else {}
    for en_txt, de_txt in de_extra.items():
        if en_txt not in en["exact"] and en_txt in tr["exact"]:
            out.append(("dash:exact:" + en_txt, "Dashboard", priority(en_txt, False, True), de_txt, en_txt, tr["exact"][en_txt]))
    # Grammar refinements of this language only (singular forms a pattern would get wrong).
    for k, v in tr["exact"].items():
        if k not in en["exact"] and k not in de_extra:
            out.append(("dash:exact:" + k, "Dashboard (singular)", 1, k, "", v))
    for i, (rx, en_txt) in enumerate(en["patterns"]):
        mine = tr["patterns"][i][1] if i < len(tr["patterns"]) and tr["patterns"][i][0] == rx else ""
        out.append(("dash:pattern:" + rx, "Dashboard (pattern)", priority(en_txt, True, True), rx, en_txt, mine))
    # Studio: English source, German table beside it.
    de = load(os.path.join(TR, "panel/de.json"))
    tr = load(os.path.join(TR, "panel/%s.json" % lang))
    for en_txt, de_txt in de["exact"].items():
        out.append(("panel:exact:" + en_txt, "Studio", priority(en_txt, False, False), de_txt, en_txt, tr["exact"].get(en_txt, "")))
    en_path = os.path.join(TR, "panel/en.json")
    en_extra = load(en_path).get("exact", {}) if os.path.exists(en_path) else {}
    for k, v in tr["exact"].items():
        if k not in de["exact"]:
            out.append(("panel:exact:" + k, "Studio" if k in en_extra else "Studio (singular)", 2, "", en_extra.get(k, k), v))
    for i, (rx, de_txt) in enumerate(de["patterns"]):
        mine = tr["patterns"][i][1] if i < len(tr["patterns"]) and tr["patterns"][i][0] == rx else ""
        out.append(("panel:pattern:" + rx, "Studio (pattern)", priority(de_txt, False, False), de_txt, rx, mine))
    # Small key tables: dashboard keys and Home Assistant's setup texts.
    for name, where in (("dashboard/%s.json", "Dashboard (keys)"), ("%s.json", "Home Assistant setup")):
        en = dict(flat(load(os.path.join(TR, name % "en"))))
        de = dict(flat(load(os.path.join(TR, name % "de")))) if os.path.exists(os.path.join(TR, name % "de")) else {}
        tr = dict(flat(load(os.path.join(TR, name % lang))))
        for k, en_txt in en.items():
            out.append(("%s:%s" % (name.split("%s")[0] or "ha", k), where, priority(en_txt, False, where.startswith("Dashboard")), de.get(k, ""), en_txt, tr.get(k, "")))
    return out


def rows_at(lang, commit):
    """rows() of an earlier commit (translation files read from git)."""
    rel = os.path.relpath(TR, ROOT)
    with tempfile.TemporaryDirectory() as tmp:
        names = subprocess.run(["git", "-C", ROOT, "ls-tree", "-r", "--name-only", commit, rel + "/"],
                               check=True, capture_output=True, text=True).stdout.split()
        for name in names:
            if not name.endswith(".json"):
                continue
            dest = os.path.join(tmp, os.path.relpath(name, rel))
            os.makedirs(os.path.dirname(dest), exist_ok=True)
            blob = subprocess.run(["git", "-C", ROOT, "show", "%s:%s" % (commit, name)], check=True, capture_output=True).stdout
            with open(dest, "wb") as f:
                f.write(blob)
        return rows(lang, tmp)


def changes(lang, old_rows):
    """Rows new or changed against old_rows, each with the reason ("New", "Source changed", "Translation changed")."""
    old = {r[0]: r for r in old_rows}
    out = []
    for r in rows(lang):
        o = old.get(r[0])
        if o is None:
            out.append((r, "New"))
        elif (o[3], o[4]) != (r[3], r[4]):
            out.append((r, "Source changed"))
        elif o[5] != r[5]:
            out.append((r, "Translation changed"))
    return out


def export(lang, path, since=None):
    if since:
        delta = changes(lang, rows_at(lang, since))
        data_rows, why = [r for r, _ in delta], {r[0]: w for r, w in delta}
    else:
        data_rows, why = rows(lang), {}
    wb = Workbook()
    info = wb.active
    info.title = "How to"
    lines = [
        ("Casora translation review: %s%s" % (lang, " – addendum: only texts new or changed since the first sheet"
                                                   if since else ""), True),
        ("", False),
        ("Thank you for helping! Each row on the sheet \"Texts\" is one text of Casora.", False),
        ("1. Read the column \"Translation\" (German and English are there for context).", False),
        ("2. In \"Correct?\" choose OK, or Wrong.", False),
        ("3. If Wrong, write the right text in \"Correction\". A short note in \"Comment\" helps.", False),
        ("4. Keep placeholders exactly as they are: {1}, {2}, {n}, {x}, %s and tags like <b>.", False),
        ("   They are replaced by values, e.g. \"Battery {1} %\" shows \"Battery 82 %\".", False),
        ("5. Texts are short on purpose: tiles and buttons are narrow.", False),
        ("6. Start with priority 1 (filter the column \"Priority\"): the dashboard texts everyone sees daily.", False),
        ("   Priority 2 is the Studio (the editor), priority 3 long help texts. Any part you check helps.", False),
        ("7. Rows marked \"Dashboard (pattern)\" show a pattern in the German column; the English column shows", False),
        ("   how it reads. {1} is the part the pattern captures (a number, a name).", False),
        ("8. Save the file and send it back. Do not change the hidden column \"Key\".", False),
        ("", False),
        ("Progress", True),
    ]
    for r, (text, bold) in enumerate(lines, start=1):
        c = info.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=14 if r == 1 else 11)
    base = len(lines) + 1
    stats = [
        ("Texts", "=COUNTA(Texts!A:A)-1"),
        ("Checked", "=COUNTIF(Texts!G:G,\"OK\")+COUNTIF(Texts!G:G,\"Wrong\")"),
        ("Marked wrong", "=COUNTIF(Texts!G:G,\"Wrong\")"),
        ("Priority 1 texts", "=COUNTIF(Texts!C:C,1)"),
        ("Priority 1 checked", "=COUNTIFS(Texts!C:C,1,Texts!G:G,\"OK\")+COUNTIFS(Texts!C:C,1,Texts!G:G,\"Wrong\")"),
    ]
    for i, (label, formula) in enumerate(stats):
        info.cell(row=base + i, column=1, value=label).font = Font(name=FONT)
        info.cell(row=base + i, column=2, value=formula).font = Font(name=FONT, bold=True)
    info.column_dimensions["A"].width = 100
    info.column_dimensions["B"].width = 12

    ws = wb.create_sheet("Texts")
    ws.append(COLS + (["Change"] if since else []))
    head = PatternFill("solid", fgColor="3A322B")
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(vertical="center")
    data = sorted(data_rows, key=lambda r: (r[2], r[1], len(r[4])))
    fill_in = PatternFill("solid", fgColor="FFF4C2")
    wrap = Alignment(wrap_text=True, vertical="top")
    for r in data:
        ws.append(list(r) + ["", "", ""] + ([why[r[0]]] if since else []))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT)
            c.alignment = wrap
        row[6].fill = fill_in
        row[7].fill = fill_in
        row[8].fill = fill_in
    # One worked example as a comment on the first row (no extra data row that import would have to skip).
    ws["H2"].comment = Comment("Example: Translation \"No informa\" marked Wrong, Correction \"Sin respuesta\".", "Casora")
    dv = DataValidation(type="list", formula1='"OK,Wrong"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("G2:G%d" % (len(data) + 1))
    widths = {"A": 4, "B": 20, "C": 8, "D": 42, "E": 42, "F": 42, "G": 11, "H": 42, "I": 30, "J": 20}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["A"].hidden = True
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = "%s1:%s%d" % ("A", "J" if since else "I", len(data) + 1)
    wb.active = 0
    wb.save(path)
    print("%s: %d texts (priority 1: %d) for %s%s" % (path, len(data), sum(1 for r in data if r[2] == 1), lang,
                                                    " since %s" % since if since else ""))


def slots(s):
    return sorted(SLOT.findall(s or ""))


def set_path(obj, path, value):
    keys = path.strip("/").split("/")
    for k in keys[:-1]:
        obj = obj[k]
    obj[keys[-1]] = value


def import_(lang, path, dry):
    src = {r[0]: r for r in rows(lang)}
    ws = load_workbook(path)["Texts"]
    head = [c.value for c in ws[1]]
    ix = {name: head.index(name) for name in COLS}
    files = {
        "dash": os.path.join(TR, "dashboard/phrases/%s.json" % lang),
        "panel": os.path.join(TR, "panel/%s.json" % lang),
        "dashboard/": os.path.join(TR, "dashboard/%s.json" % lang),
        "ha": os.path.join(TR, "%s.json" % lang),
    }
    data = {k: load(p) for k, p in files.items()}
    cell = lambda row, name: "" if row[ix[name]] is None else str(row[ix[name]])  # noqa: E731
    done, skipped, conflicts, same, open_ = 0, [], [], 0, []
    for row in ws.iter_rows(min_row=2, values_only=True):
        key, verdict, fix = row[ix["Key"]], cell(row, "Correct?").strip().lower(), cell(row, "Correction").strip()
        if not key:
            continue
        if not fix:
            if verdict == "wrong":
                open_.append(key)
            continue
        if verdict == "ok":
            continue
        if key not in src:
            skipped.append((key, "text no longer in Casora"))
            continue
        now = src[key]
        # Only where the row still shows what the reviewer saw: source (German, English) and translation.
        if (cell(row, "German"), cell(row, "English")) != (now[3], now[4]):
            conflicts.append((key, "source changed since the sheet", cell(row, "Translation"), now[5], fix))
            continue
        if now[5] == fix:
            same += 1
            continue
        if cell(row, "Translation") != now[5]:
            conflicts.append((key, "translation changed since the sheet", cell(row, "Translation"), now[5], fix))
            continue
        ref = now[4] if not key.startswith("panel:pattern:") else now[3]
        if slots(fix) != slots(ref):
            skipped.append((key, "placeholders differ: %s vs %s" % (slots(fix), slots(ref))))
            continue
        kind, _, rest = key.partition(":")
        if kind in ("dash", "panel"):
            part, _, k = rest.partition(":")
            table = data[kind]
            if part == "exact":
                table["exact"][k] = fix
            else:
                for p in table["patterns"]:
                    if p[0] == k:
                        p[1] = fix
        else:
            set_path(data["dashboard/" if kind == "dashboard/" else "ha"], rest, fix)
        done += 1
    for key, why in skipped:
        print("skipped %s: %s" % (key[:80], why))
    for key, why, saw, now_txt, fix in conflicts:
        print("CONFLICT %s: %s\n  sheet: %r\n  now:   %r\n  fix:   %r" % (key[:80], why, saw, now_txt, fix))
    for key in open_:
        print("marked wrong without a correction: %s" % key[:80])
    print("%d corrections, %d already in, %d conflicts (kept the current text), %d skipped, %d wrong without correction%s"
          % (done, same, len(conflicts), len(skipped), len(open_), " (dry run)" if dry else ""))
    if dry:
        return {"done": done, "conflicts": conflicts, "skipped": skipped, "same": same, "open": open_}
    indent = {"dash": 1, "panel": 2, "dashboard/": 2, "ha": 2}
    for k, p in files.items():
        if data[k] == load(p):
            continue
        with open(p, "w", encoding="utf-8") as f:
            f.write(json.dumps(data[k], ensure_ascii=False, indent=indent[k]) + "\n")
    if done:
        print("next: python3 tools/build-i18n.py && python3 tools/build-panel-i18n.py && python3 tools/i18ncheck.py")
    return {"done": done, "conflicts": conflicts, "skipped": skipped, "same": same, "open": open_}


if __name__ == "__main__":
    a = sys.argv[1:]
    if len(a) >= 2 and a[0] == "export":
        since = a[a.index("--since") + 1] if "--since" in a else None
        pos = [x for i, x in enumerate(a) if not x.startswith("--") and (i == 0 or a[i - 1] != "--since")]
        export(a[1], pos[2] if len(pos) > 2 else "casora-%s-review.xlsx" % a[1], since)
    elif len(a) >= 3 and a[0] == "import":
        import_(a[1], a[2], "--dry-run" in a)
    else:
        print(__doc__)
        sys.exit(2)
